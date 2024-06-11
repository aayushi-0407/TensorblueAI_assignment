# file_reader/utils.py

import PyPDF2
import pdfplumber
import pandas as pd
import openpyxl
import docx

def read_pdf(file_path):
    content = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            content += page.extract_text() + "\n"
    return content

def read_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(row)
    return data

def read_csv(file_path):
    df = pd.read_csv(file_path)
    return df.to_dict(orient='records')

def read_docx(file_path):
    doc = docx.Document(file_path)
    content = ""
    for para in doc.paragraphs:
        content += para.text + "\n"
    return content

# file_reader/utils.py

